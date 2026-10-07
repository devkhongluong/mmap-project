"""
Script tạo file Excel template mẫu để import lộ trình vào MMAP.
Chạy: python create_template.py

Yêu cầu: pip install openpyxl
"""

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Cài đặt openpyxl: pip install openpyxl")
    exit(1)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Lộ trình"

# ── Styles ──────────────────────────────────────────────────────────────
header_font   = Font(bold=True, color="FFFFFF", size=11)
header_fill   = PatternFill("solid", fgColor="6366F1")   # Indigo
example_fill  = PatternFill("solid", fgColor="EEF2FF")   # Light indigo
center        = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

# ── Headers ─────────────────────────────────────────────────────────────
headers = [
    "Phase", "Week", "Day", "Title",
    "Check 1", "Check 2", "Check 3", "Check 4"
]
col_widths = [25, 15, 6, 35, 50, 50, 50, 50]

for col, (h, w) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font      = header_font
    cell.fill      = header_fill
    cell.alignment = center
    cell.border    = thin_border
    ws.column_dimensions[get_column_letter(col)].width = w

ws.row_dimensions[1].height = 30

# ── Example data ─────────────────────────────────────────────────────────
sample_rows = [
    # Phase 1 — Java Core
    ("Phase 1: Java Core", "Tuần 1", 1, "Nền tảng OOP & Class",
     "Hiểu class, object, constructor và instance variables",
     "Phân biệt static vs instance method",
     "Viết class BankAccount với deposit/withdraw",
     ""),
    ("Phase 1: Java Core", "Tuần 1", 2, "Kế thừa & Encapsulation",
     "Dùng extends và hiểu inheritance chain",
     "Hiểu private/protected/public access modifiers",
     "Implement getter/setter đúng chuẩn JavaBean",
     "Code bài tập: Shape → Circle, Rectangle"),
    ("Phase 1: Java Core", "Tuần 1", 3, "Đa hình & Interface",
     "Hiểu Polymorphism và runtime method dispatch",
     "Phân biệt Abstract Class vs Interface",
     "Implement Comparable để sort Custom Object",
     "Code: hệ thống quản lý động vật đa hình"),
    ("Phase 1: Java Core", "Tuần 1", 4, "Collections Framework",
     "Dùng ArrayList, LinkedList, HashMap thành thạo",
     "Hiểu Big-O của các thao tác CRUD trong mỗi Collection",
     "So sánh List vs Set vs Map khi nào dùng cái nào",
     ""),
    ("Phase 1: Java Core", "Tuần 1", 5, "Stream API & Lambda",
     "Viết lambda expression thay vì anonymous class",
     "Dùng stream().filter().map().collect() cơ bản",
     "Xử lý Optional để tránh NullPointerException",
     ""),

    # Phase 2 — Spring Boot
    ("Phase 2: Spring Boot", "Tuần 5", 29, "Spring Boot Setup",
     "Tạo project Spring Boot 3 với Spring Initializr",
     "Hiểu @SpringBootApplication và auto-configuration",
     "Chạy Hello World API với @RestController",
     ""),
    ("Phase 2: Spring Boot", "Tuần 5", 30, "REST API & HTTP Methods",
     "Implement CRUD API đầy đủ cho một Entity",
     "Dùng @GetMapping @PostMapping @PutMapping @DeleteMapping",
     "Test với Postman — kiểm tra status codes",
     "Xử lý @PathVariable và @RequestParam"),
]

for row_idx, row_data in enumerate(sample_rows, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill      = example_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border    = thin_border
    ws.row_dimensions[row_idx].height = 45

# ── Instructions sheet ───────────────────────────────────────────────────
ws2 = wb.create_sheet("Hướng dẫn")
instructions = [
    ("HƯỚNG DẪN IMPORT LỘ TRÌNH VÀO MMAP", ),
    ("",),
    ("Cột A — Phase:", "Tên giai đoạn học. VD: 'Phase 1: Java Core'"),
    ("Cột B — Week:", "Tên tuần học (không bắt buộc). VD: 'Tuần 1'"),
    ("Cột C — Day:", "Số thứ tự ngày học (1, 2, 3...). BẮT BUỘC, phải là số nguyên"),
    ("Cột D — Title:", "Tên ngày học, hiển thị trên Tree Map. BẮT BUỘC"),
    ("Cột E, F, G, H...:", "Nội dung các mục checkbox trong ngày (không bắt buộc)"),
    ("",),
    ("LƯU Ý:", ""),
    ("- Sheet phải tên là 'Lộ trình' hoặc là sheet đầu tiên trong file",),
    ("- Hàng đầu tiên là header — sẽ bị bỏ qua khi import",),
    ("- Hàng trống sẽ bị bỏ qua tự động",),
    ("- Tối đa 10MB per file",),
    ("- Chỉ hỗ trợ định dạng .xlsx",),
]

ws2.column_dimensions["A"].width = 25
ws2.column_dimensions["B"].width = 70

title_cell = ws2.cell(row=1, column=1, value="HƯỚNG DẪN IMPORT LỘ TRÌNH VÀO MMAP")
title_cell.font = Font(bold=True, size=14, color="6366F1")

for row_idx, row_data in enumerate(instructions[1:], 2):
    for col_idx, value in enumerate(row_data, 1):
        ws2.cell(row=row_idx, column=col_idx, value=value)

# ── Save ─────────────────────────────────────────────────────────────────
output_file = "mmap_template.xlsx"
wb.save(output_file)
print(f"✅ Đã tạo file template: {output_file}")
print(f"   → Upload file này lên MMAP tại: POST /api/maps/import")
print(f"   → Hoặc dùng tính năng Import trong Dashboard")
